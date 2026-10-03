# -*- coding: utf-8 -*-
"""
邮件群发助手（图形界面版）

功能：
- 读取 Excel(.xlsx/.xls) 或 CSV 收件人数据
- 邮件主题和正文(HTML)完全自定义，任意列可用占位符 {列名} 插入
- 日期列可用 {日期:%Y年%m月%d日} 自定义格式
- 支持附件、发送预览、测试发送、失败名单导出、配置记忆
"""

import os
import re
import csv
import sys
import json
import time
import queue
import threading
import smtplib
import webbrowser
import tempfile
import traceback
from datetime import datetime, date, time as dtime, timedelta

from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
from email.header import Header
from email.utils import formataddr

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext, font as tkfont

try:
    import openpyxl  # 读取 .xlsx
except ImportError:
    openpyxl = None

try:
    import xlrd  # 读取 .xls
except ImportError:
    xlrd = None


# 占位符语法：{列名} 或 {列名:格式}，例如 {姓名}、{日期:%Y年%m月%d日}
PLACEHOLDER_RE = re.compile(r"\{([^{}\n:]+)(?::([^{}\n]+))?\}")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")

SMTP_PRESETS = {
    "163.com": ("smtp.163.com", 465, "SSL"),
    "126.com": ("smtp.126.com", 465, "SSL"),
    "qq.com": ("smtp.qq.com", 465, "SSL"),
    "foxmail.com": ("smtp.qq.com", 465, "SSL"),
    "gmail.com": ("smtp.gmail.com", 465, "SSL"),
    "outlook.com": ("smtp.office365.com", 587, "STARTTLS"),
    "hotmail.com": ("smtp.office365.com", 587, "STARTTLS"),
    "aliyun.com": ("smtp.aliyun.com", 465, "SSL"),
    "sina.com": ("smtp.sina.com", 465, "SSL"),
    "sohu.com": ("smtp.sohu.com", 465, "SSL"),
}

DEFAULT_SUBJECT = "xx市xxx有限公司面试通知"
DEFAULT_BODY = """<p>{姓名}你好！</p>
<p>你已通过xx市xxx有限公司2026届校园招聘xxx部xxx岗资格审查，现将面试有关事项通知如下：</p>
<p>一、面试日期：{日期:%Y-%m-%d}</p>
<p>二、面试时段：{时段}</p>
<p>三、面试时间：{时间}</p>
<p>四、面试顺序：{顺序}</p>
<p>五、面试注意事项：{面试注意事项}</p>
<p>请回复已收到邮件，确认参加面试。</p>"""

SAMPLE_ROW = {
    "姓名": "张三",
    "邮箱": "test@example.com",
    "日期": datetime.now(),
    "时间": dtime(9, 30),
    "时段": "上午",
    "顺序": 1,
    "面试注意事项": "请携带身份证与学生证，提前15分钟到场。",
}


def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def config_path():
    return os.path.join(app_dir(), "mail_config.json")


# ---------------------------------------------------------------- 数据读取

def read_table(path):
    """读取 Excel/CSV，返回 (headers, rows)，rows 为 list[dict]"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext == ".xls":
        return _read_xls(path)
    if ext == ".csv":
        return _read_csv(path)
    raise ValueError("不支持的文件类型：%s（支持 .xlsx / .xls / .csv）" % ext)


def _clean_headers(raw):
    headers = []
    for h in raw:
        h = "" if h is None else str(h).strip()
        headers.append(h if h else "未命名列")
    return headers


def _read_xlsx(path):
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl 库，无法读取 .xlsx")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = None
    rows = []
    for r in rows_iter:
        if headers is None:
            headers = _clean_headers(r)
            continue
        rows.append({h: v for h, v in zip(headers, r)})
    wb.close()
    return headers, rows


def _read_xls(path):
    if xlrd is None:
        raise RuntimeError("缺少 xlrd 库，无法读取 .xls")
    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    headers = _clean_headers([sheet.cell_value(0, c) for c in range(sheet.ncols)])
    rows = []
    for r in range(1, sheet.nrows):
        row = {}
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            v = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    v = xlrd.xldate_as_datetime(v, book.datemode)
                except Exception:
                    pass
            row[headers[c]] = v
        rows.append(row)
    return headers, rows


def _read_csv(path):
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                reader = list(csv.reader(f))
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        raise ValueError("无法识别 CSV 文件编码")
    if not reader:
        return [], []
    headers = _clean_headers(reader[0])
    rows = [{h: v for h, v in zip(headers, r)} for r in reader[1:] if any(x.strip() for x in r)]
    return headers, rows


# ---------------------------------------------------------------- 模板渲染

def fmt_value(v, fmt=None):
    if v is None:
        return ""
    if isinstance(v, datetime):
        # xlrd 读取的纯时间单元格会带 1899 年的基准日期，此时按时间格式输出
        if (v.hour or v.minute or v.second) and v.year < 1901:
            return v.strftime(fmt or "%H:%M")
        return v.strftime(fmt or "%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime(fmt or "%Y-%m-%d")
    if isinstance(v, dtime):
        return v.strftime(fmt or "%H:%M")
    if isinstance(v, timedelta):
        return str(v)
    if fmt is not None:
        try:
            return format(v, fmt)
        except Exception:
            return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def render_template(template, row):
    def sub(m):
        name = m.group(1).strip()
        fmt = m.group(2)
        if name in row:
            return fmt_value(row[name], fmt)
        return m.group(0)  # 未识别的占位符原样保留（<style> 块中的花括号也不会被误改）

    return PLACEHOLDER_RE.sub(sub, template)


def find_placeholders(template):
    return [m.group(1).strip() for m in PLACEHOLDER_RE.finditer(template)]


def wrap_html(body):
    return (
        '<!DOCTYPE html><html><head><meta charset="utf-8"></head>'
        '<body style="font-family:Microsoft YaHei,Segoe UI,Arial,sans-serif;">'
        + body + "</body></html>"
    )


# ---------------------------------------------------------------- 邮件发送

def connect_smtp(host, port, encryption):
    if encryption == "SSL":
        server = smtplib.SMTP_SSL(host, int(port), timeout=30)
    elif encryption == "STARTTLS":
        server = smtplib.SMTP(host, int(port), timeout=30)
        server.ehlo()
        server.starttls()
        server.ehlo()
    else:
        server = smtplib.SMTP(host, int(port), timeout=30)
    return server


def build_message(sender, sender_name, to_list, subject, html, attachments):
    msg = MIMEMultipart()
    if sender_name:
        msg["From"] = formataddr((str(Header(sender_name, "utf-8")), sender))
    else:
        msg["From"] = sender
    msg["To"] = ", ".join(to_list)
    msg["Subject"] = Header(subject, "utf-8")
    msg.attach(MIMEText(html, "html", "utf-8"))
    for path in attachments:
        with open(path, "rb") as f:
            data = f.read()
        part = MIMEBase("application", "octet-stream")
        part.set_payload(data)
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", "attachment",
                        filename=("utf-8", "", os.path.basename(path)))
        msg.attach(part)
    return msg


def login(server, sender, password):
    try:
        if password:
            server.login(sender, password)
    except smtplib.SMTPAuthenticationError:
        raise RuntimeError(
            "SMTP 登录失败：请确认已在该邮箱网页设置中开启 SMTP 服务，"
            "并使用「授权码」（而非邮箱登录密码）"
        )


def send_one(server, sender, sender_name, to_list, subject, html, attachments):
    msg = build_message(sender, sender_name, to_list, subject, html, attachments)
    server.sendmail(sender, to_list, msg.as_string())


# ---------------------------------------------------------------- 界面

class MailApp:

    def __init__(self, root):
        self.root = root
        root.title("邮件群发助手")
        # 按系统 DPI 缩放窗口，避免高分屏下内容溢出
        scale = root.winfo_fpixels("1i") / 96.0
        root.geometry("%dx%d" % (int(980 * scale), int(880 * scale)))
        root.minsize(int(920 * scale), int(820 * scale))

        self.log_queue = queue.Queue()
        self.stop_event = threading.Event()
        self.send_thread = None
        self.insert_widget = None
        self.headers = []
        self.rows = []
        self.recipient_col = None
        self.attachments = []

        self._build_vars()
        self._build_ui()
        self._bind_keys()
        self._load_config()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.root.after(80, self._poll_log)
        self._log("欢迎使用邮件群发助手。先选择 Excel/CSV 收件数据，再用 {列名} 占位符编写模板。")

    # ---------- 控件 ----------

    def _build_vars(self):
        self.var_file = tk.StringVar()
        self.var_sender = tk.StringVar()
        self.var_password = tk.StringVar()
        self.var_nickname = tk.StringVar()
        self.var_host = tk.StringVar()
        self.var_port = tk.StringVar(value="465")
        self.var_enc = tk.StringVar(value="SSL")
        self.var_interval = tk.StringVar(value="2")
        self.var_remember = tk.BooleanVar(value=False)
        self.var_status = tk.StringVar(value="尚未加载数据")

    def _build_ui(self):
        pad = dict(padx=8, pady=4)
        style_font = ("Microsoft YaHei UI", 10)

        # ---- 收件数据 ----
        f_data = ttk.LabelFrame(self.root, text=" 1. 收件数据（Excel / CSV） ")
        f_data.pack(fill="x", **pad)
        row1 = ttk.Frame(f_data); row1.pack(fill="x", padx=6, pady=4)
        ttk.Label(row1, text="数据文件：").pack(side="left")
        ttk.Entry(row1, textvariable=self.var_file, state="readonly").pack(
            side="left", fill="x", expand=True, padx=4)
        ttk.Button(row1, text="选择文件", command=self._choose_file).pack(side="left")
        row2 = ttk.Frame(f_data); row2.pack(fill="x", padx=6, pady=(0, 6))
        ttk.Label(row2, text="可插入的列：").pack(side="left")
        self.chip_frame = ttk.Frame(row2)
        self.chip_frame.pack(side="left", fill="x", expand=True)
        ttk.Label(f_data, textvariable=self.var_status, foreground="#555").pack(
            anchor="w", padx=8, pady=(0, 4))

        # ---- 发件设置 ----
        f_sender = ttk.LabelFrame(self.root, text=" 2. 发件邮箱设置 ")
        f_sender.pack(fill="x", **pad)
        r1 = ttk.Frame(f_sender); r1.pack(fill="x", padx=6, pady=4)
        ttk.Label(r1, text="发件邮箱：").pack(side="left")
        ttk.Entry(r1, textvariable=self.var_sender, width=24).pack(side="left", padx=(0, 10))
        ttk.Label(r1, text="授权码/密码：").pack(side="left")
        ttk.Entry(r1, textvariable=self.var_password, show="*", width=20).pack(side="left", padx=(0, 10))
        ttk.Checkbutton(r1, text="记住授权码", variable=self.var_remember).pack(side="left")
        r2 = ttk.Frame(f_sender); r2.pack(fill="x", padx=6, pady=(0, 6))
        ttk.Label(r2, text="SMTP服务器：").pack(side="left")
        ttk.Entry(r2, textvariable=self.var_host, width=20).pack(side="left", padx=(0, 10))
        ttk.Label(r2, text="端口：").pack(side="left")
        ttk.Entry(r2, textvariable=self.var_port, width=6).pack(side="left", padx=(0, 10))
        ttk.Label(r2, text="加密方式：").pack(side="left")
        ttk.Combobox(r2, textvariable=self.var_enc, values=["SSL", "STARTTLS", "无"],
                     width=9, state="readonly").pack(side="left", padx=(0, 10))
        ttk.Label(r2, text="发件人昵称(可选)：").pack(side="left")
        ttk.Entry(r2, textvariable=self.var_nickname, width=12).pack(side="left")

        # ---- 邮件模板 ----
        f_tpl = ttk.LabelFrame(self.root, text=" 3. 邮件模板（占位符写法：{列名}，日期可用 {日期:%Y年%m月%d日}） ")
        f_tpl.pack(fill="both", expand=True, **pad)
        r3 = ttk.Frame(f_tpl); r3.pack(fill="x", padx=6, pady=4)
        ttk.Label(r3, text="邮件主题：").pack(side="left")
        self.ent_subject = ttk.Entry(r3, width=50)
        self.ent_subject.insert(0, DEFAULT_SUBJECT)
        self.ent_subject.pack(side="left", fill="x", expand=True)
        ttk.Label(f_tpl, text="正文模板（HTML，点上方“可插入的列”按钮可快速插入占位符）：").pack(
            anchor="w", padx=8)
        self.txt_body = scrolledtext.ScrolledText(f_tpl, height=9, font=style_font, wrap="word",
                                                  undo=True)
        self.txt_body.insert("1.0", DEFAULT_BODY)
        self.txt_body.pack(fill="both", expand=True, padx=6, pady=2)
        r4 = ttk.Frame(f_tpl); r4.pack(fill="x", padx=6, pady=(2, 6))
        ttk.Label(r4, text="附件：").pack(side="left")
        self.lst_att = tk.Listbox(r4, height=3)
        self.lst_att.pack(side="left", fill="both", expand=True, padx=4)
        att_btns = ttk.Frame(r4); att_btns.pack(side="left", fill="y")
        for text, cmd in (("添加", self._add_attachment), ("移除", self._remove_attachment),
                          ("清空", self._clear_attachments)):
            ttk.Button(att_btns, text=text, width=6, command=cmd).pack(pady=1)

        # ---- 发送控制 ----
        f_ctrl = ttk.LabelFrame(self.root, text=" 4. 发送 ")
        f_ctrl.pack(fill="x", **pad)
        r5 = ttk.Frame(f_ctrl); r5.pack(fill="x", padx=6, pady=4)
        ttk.Label(r5, text="发送间隔(秒)：").pack(side="left")
        ttk.Spinbox(r5, textvariable=self.var_interval, from_=0, to=60, width=5).pack(side="left", padx=(0, 14))
        self.btn_preview = ttk.Button(r5, text="预览效果", command=self._preview)
        self.btn_preview.pack(side="left", padx=4)
        self.btn_test = ttk.Button(r5, text="测试发送(发给自己)", command=self._test_send)
        self.btn_test.pack(side="left", padx=4)
        self.btn_start = ttk.Button(r5, text="开始发送", command=self._start_send)
        self.btn_start.pack(side="left", padx=4)
        self.btn_stop = ttk.Button(r5, text="停止", command=self._stop_send, state="disabled")
        self.btn_stop.pack(side="left", padx=4)
        self.progress = ttk.Progressbar(f_ctrl, mode="determinate")
        self.progress.pack(fill="x", padx=8, pady=(0, 6))

        # ---- 日志 ----
        f_log = ttk.LabelFrame(self.root, text=" 发送日志 ")
        f_log.pack(fill="both", expand=True, **pad)
        self.txt_log = scrolledtext.ScrolledText(f_log, height=7, font=("Consolas", 9), state="disabled")
        self.txt_log.pack(fill="both", expand=True, padx=6, pady=4)
        self.txt_log.tag_configure("err", foreground="#c0392b")
        self.txt_log.tag_configure("ok", foreground="#2e7d32")

    def _bind_keys(self):
        for w in (self.ent_subject, self.txt_body):
            w.bind("<FocusIn>", lambda e: setattr(self, "insert_widget", e.widget))
        self.var_sender.trace_add("write", self._auto_smtp)

    # ---------- 配置持久化 ----------

    def _load_config(self):
        try:
            with open(config_path(), "r", encoding="utf-8") as f:
                cfg = json.load(f)
        except Exception:
            return
        self.var_sender.set(cfg.get("sender", ""))
        self.var_nickname.set(cfg.get("nickname", ""))
        self.var_host.set(cfg.get("host", ""))
        self.var_port.set(cfg.get("port", "465"))
        self.var_enc.set(cfg.get("enc", "SSL"))
        self.var_interval.set(str(cfg.get("interval", 2)))
        if cfg.get("remember_pwd"):
            self.var_password.set(cfg.get("password", ""))
            self.var_remember.set(True)
        subject = cfg.get("subject")
        body = cfg.get("body")
        if subject is not None:
            self.ent_subject.delete(0, "end"); self.ent_subject.insert(0, subject)
        if body:
            self.txt_body.delete("1.0", "end"); self.txt_body.insert("1.0", body)
        self._auto_smtp()

    def _save_config(self):
        cfg = {
            "sender": self.var_sender.get().strip(),
            "nickname": self.var_nickname.get().strip(),
            "host": self.var_host.get().strip(),
            "port": self.var_port.get().strip(),
            "enc": self.var_enc.get(),
            "interval": self._get_interval(),
            "subject": self.ent_subject.get(),
            "body": self.txt_body.get("1.0", "end").rstrip("\n"),
            "remember_pwd": self.var_remember.get(),
        }
        if self.var_remember.get():
            cfg["password"] = self.var_password.get()
        try:
            with open(config_path(), "w", encoding="utf-8") as f:
                json.dump(cfg, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    # ---------- 工具 ----------

    def _log(self, msg, tag=None):
        self.log_queue.put(("log", msg, tag))

    def _now(self):
        return datetime.now().strftime("%H:%M:%S")

    def _get_interval(self):
        try:
            return max(0.0, float(self.var_interval.get()))
        except ValueError:
            return 2.0

    def _collect_settings(self):
        return {
            "sender": self.var_sender.get().strip(),
            "password": self.var_password.get(),
            "nickname": self.var_nickname.get().strip(),
            "host": self.var_host.get().strip(),
            "port": self.var_port.get().strip(),
            "enc": self.var_enc.get(),
            "subject_tpl": self.ent_subject.get(),
            "body_tpl": self.txt_body.get("1.0", "end").rstrip("\n"),
            "attachments": list(self.attachments),
        }

    def _auto_smtp(self, *args):
        addr = self.var_sender.get().strip()
        m = EMAIL_RE.search(addr)
        if not m:
            return
        domain = m.group(0).split("@")[1].lower()
        preset = SMTP_PRESETS.get(domain)
        if preset and not self.var_host.get().strip():
            self.var_host.set(preset[0])
            self.var_port.set(str(preset[1]))
            self.var_enc.set(preset[2])

    def _render_row(self, cfg, row):
        subject = render_template(cfg["subject_tpl"], row)
        body = wrap_html(render_template(cfg["body_tpl"], row))
        return subject, body

    def _recipient_of(self, row):
        raw = row.get(self.recipient_col, "")
        return EMAIL_RE.findall(fmt_value(raw))

    # ---------- 数据加载 ----------

    def _choose_file(self):
        path = filedialog.askopenfilename(
            title="选择收件数据文件",
            filetypes=[("表格文件", "*.xlsx *.xlsm *.xls *.csv"), ("所有文件", "*.*")])
        if not path:
            return
        try:
            self._load_table(path)
        except Exception as e:
            messagebox.showerror("读取失败", str(e))

    def _load_table(self, path):
        headers, rows = read_table(path)
        if not rows:
            raise ValueError("文件中没有数据行")
        self.headers = headers
        self.rows = rows
        self.var_file.set(path)

        self.recipient_col = None
        for h in headers:
            if any(k in h.lower() for k in ("邮箱", "邮件", "email", "mail")):
                self.recipient_col = h
                break

        for w in self.chip_frame.winfo_children():
            w.destroy()
        per_row = 7
        for i, h in enumerate(headers):
            # 不设固定宽度，让按钮按文字自适应（中文比 ttk 宽度基准宽，固定值会截断）
            b = ttk.Button(self.chip_frame, text="{" + h + "}",
                           command=lambda n=h: self._insert_placeholder(n))
            b.grid(row=i // per_row, column=i % per_row, padx=2, pady=1, sticky="w")

        if self.recipient_col is None:
            self.var_status.set("共 %d 条记录，但未找到“邮箱”列（列名需包含“邮箱”或 email 字样）" % len(rows))
            messagebox.showwarning("缺少邮箱列",
                                   "没有找到收件邮箱列。请确保表格中有一列列名包含“邮箱”（如：邮箱、收件邮箱、Email）。")
        else:
            self.var_status.set("共 %d 条记录，收件邮箱列：「%s」" % (len(rows), self.recipient_col))
        self._log("已加载 %s：%d 条记录，列：%s" % (os.path.basename(path), len(rows), "、".join(headers)))

    def _insert_placeholder(self, name):
        w = self.insert_widget or self.txt_body
        text = "{" + name + "}"
        if isinstance(w, tk.Text):
            w.insert("insert", text)
        else:
            w.insert("insert", text)
        w.focus_set()

    # ---------- 附件 ----------

    def _add_attachment(self):
        paths = filedialog.askopenfilenames(title="添加附件")
        for p in paths:
            if p not in self.attachments:
                self.attachments.append(p)
                self.lst_att.insert("end", p)

    def _remove_attachment(self):
        sel = self.lst_att.curselection()
        if not sel:
            return
        for i in reversed(sel):
            del self.attachments[i]
            self.lst_att.delete(i)

    def _clear_attachments(self):
        self.attachments.clear()
        self.lst_att.delete(0, "end")

    # ---------- 预览 / 测试 ----------

    def _sample_row(self):
        if self.rows:
            return self.rows[0]
        return dict(SAMPLE_ROW)

    def _preview(self):
        cfg = self._collect_settings()
        subject, body = self._render_row(cfg, self._sample_row())
        fd, path = tempfile.mkstemp(suffix=".html", prefix="mail_preview_")
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write("<!--主题：%s-->\n" % subject)
            f.write('<div style="border:1px solid #ccc;padding:6px 10px;margin-bottom:12px;'
                    'background:#f7f7f7;font-size:13px;">主题：%s</div>\n' % subject)
            f.write(body)
        webbrowser.open("file:///" + path.replace("\\", "/"))
        self._log("已生成预览（使用第 1 条数据）：%s" % path)

    def _test_send(self):
        cfg = self._collect_settings()
        if not cfg["sender"]:
            messagebox.showwarning("提示", "请先填写发件邮箱")
            return
        to = cfg["sender"]
        row = dict(self._sample_row())
        row[self.recipient_col or "邮箱"] = to
        self._set_running(True)
        self._log("正在发送测试邮件至 %s ..." % to)
        threading.Thread(target=self._worker, args=(cfg, [row], 0, True), daemon=True).start()

    # ---------- 批量发送 ----------

    def _start_send(self):
        if self.send_thread and self.send_thread.is_alive():
            return
        cfg = self._collect_settings()

        # 每次发送前重新读取文件，保证数据是最新的
        if not self.var_file.get():
            messagebox.showwarning("提示", "请先选择 Excel/CSV 数据文件")
            return
        try:
            headers, rows = read_table(self.var_file.get())
        except Exception as e:
            messagebox.showerror("读取失败", str(e))
            return
        self.headers, self.rows = headers, rows
        if self.recipient_col is None or self.recipient_col not in headers:
            for h in headers:
                if any(k in h.lower() for k in ("邮箱", "邮件", "email", "mail")):
                    self.recipient_col = h
                    break
        if self.recipient_col is None:
            messagebox.showerror("错误", "表格中未找到收件邮箱列（列名需包含“邮箱”或 email）")
            return
        if not rows:
            messagebox.showerror("错误", "表格中没有数据行")
            return
        if not cfg["sender"] or "@" not in cfg["sender"]:
            messagebox.showwarning("提示", "请填写正确的发件邮箱")
            return
        if not cfg["host"]:
            messagebox.showwarning("提示", "请填写 SMTP 服务器地址（如 smtp.163.com）")
            return
        if not cfg["subject_tpl"].strip():
            messagebox.showwarning("提示", "邮件主题不能为空")
            return
        if not cfg["password"] and not messagebox.askyesno(
                "提示", "未填写授权码/密码，多数邮箱需要授权码才能登录发送。仍要继续吗？"):
            return
        for p in cfg["attachments"]:
            if not os.path.isfile(p):
                messagebox.showerror("错误", "附件不存在：%s" % p)
                return

        unknown = sorted({n for n in find_placeholders(cfg["subject_tpl"] + cfg["body_tpl"])
                          if n not in headers})
        if unknown and not messagebox.askyesno(
                "占位符检查",
                "以下占位符在表格列中不存在：\n\n%s\n\n"
                "点击“是”继续发送（未识别的占位符将原样保留），点击“否”返回修改。"
                % "\n".join("{%s}" % u for u in unknown)):
            return

        self._save_config()
        self._set_running(True)
        self.progress.config(maximum=len(rows), value=0)
        self._log("========== 开始发送，共 %d 条 ==========" % len(rows))
        self.stop_event.clear()
        self.send_thread = threading.Thread(
            target=self._worker, args=(cfg, rows, self._get_interval(), False), daemon=True)
        self.send_thread.start()

    def _stop_send(self):
        self.stop_event.set()
        self._log("正在停止（当前邮件发完即停）...")

    def _set_running(self, running):
        state_start = "disabled" if running else "normal"
        state_stop = "normal" if running else "disabled"
        for b in (self.btn_start, self.btn_test, self.btn_preview):
            b.config(state=state_start)
        self.btn_stop.config(state=state_stop)

    def _worker(self, cfg, rows, interval, is_test):
        q = self.log_queue
        ok, failures = 0, []
        server = None
        total = len(rows)
        try:
            q.put(("log", "连接 %s:%s (%s) ..." % (cfg["host"], cfg["port"], cfg["enc"]), None))
            server = connect_smtp(cfg["host"], cfg["port"], cfg["enc"])
            login(server, cfg["sender"], cfg["password"])
            q.put(("log", "登录成功，开始发送", None))

            for i, row in enumerate(rows, 1):
                if self.stop_event.is_set():
                    q.put(("log", "已手动停止。", None))
                    break
                to_list = self._recipient_of(row)
                if not to_list:
                    failures.append((i, "", "邮箱为空或格式不正确"))
                    q.put(("log", "[%d/%d] ✗ 第%d行收件邮箱无效" % (i, total, i), "err"))
                    q.put(("progress", i))
                    continue
                subject, body = self._render_row(cfg, row)
                try:
                    send_one(server, cfg["sender"], cfg["nickname"], to_list, subject, body,
                             cfg["attachments"])
                    ok += 1
                    q.put(("log", "[%d/%d] ✓ 已发送：%s" % (i, total, ", ".join(to_list)), "ok"))
                except Exception as e:
                    failures.append((i, ", ".join(to_list), str(e)))
                    q.put(("log", "[%d/%d] ✗ 发送失败：%s — %s" % (i, total, ", ".join(to_list), e), "err"))
                q.put(("progress", i))
                if i < total and interval > 0 and not is_test:
                    for _ in range(int(interval * 10)):
                        if self.stop_event.is_set():
                            break
                        time.sleep(0.1)
            if not self.stop_event.is_set() and not is_test:
                q.put(("log", "========== 发送完成：成功 %d / 共 %d，失败 %d ==========" % (ok, total, len(failures)), None))
        except Exception as e:
            q.put(("log", "发送中断：%s" % e, "err"))
            q.put(("log", "（已成功发送 %d 封，可修正后重新发送剩余记录）" % ok, None))
        finally:
            if server is not None:
                try:
                    server.quit()
                except Exception:
                    pass
            q.put(("done", (ok, total, failures, is_test)))

    # ---------- 日志轮询 ----------

    def _poll_log(self):
        try:
            while True:
                item = self.log_queue.get_nowait()
                kind = item[0]
                if kind == "log":
                    _, msg, tag = item
                    self.txt_log.config(state="normal")
                    self.txt_log.insert("end", "[%s] %s\n" % (self._now(), msg), tag or ())
                    self.txt_log.see("end")
                    self.txt_log.config(state="disabled")
                elif kind == "progress":
                    self.progress.config(value=item[1])
                elif kind == "done":
                    self._on_send_done(*item[1])
        except queue.Empty:
            pass
        self.root.after(80, self._poll_log)

    def _on_send_done(self, ok, total, failures, is_test):
        self._set_running(False)
        if is_test:
            if failures:
                messagebox.showerror("测试失败", "测试邮件发送失败：\n%s" % failures[0][2])
            else:
                messagebox.showinfo("测试成功", "测试邮件已发送至你的发件邮箱，请查收确认格式无误。")
            return
        if failures:
            self._log("以下记录发送失败：", "err")
            for i, to, err in failures:
                self._log("  第%d行 %s — %s" % (i, to, err), "err")
            if messagebox.askyesno("发送完成",
                                   "成功 %d / 共 %d，失败 %d。\n是否将失败名单导出为 CSV？"
                                   % (ok, total, len(failures))):
                self._export_failures(failures)
        else:
            messagebox.showinfo("发送完成", "全部发送成功：%d / %d" % (ok, total))

    def _export_failures(self, failures):
        path = filedialog.asksaveasfilename(
            title="保存失败名单", defaultextension=".csv",
            initialfile="发送失败名单_%s.csv" % datetime.now().strftime("%Y%m%d_%H%M%S"),
            filetypes=[("CSV 文件", "*.csv")])
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(["行号", "收件邮箱", "错误信息"])
                for i, to, err in failures:
                    w.writerow([i, to, err])
            self._log("失败名单已保存：%s" % path, "ok")
        except Exception as e:
            messagebox.showerror("保存失败", str(e))

    # ---------- 退出 ----------

    def _on_close(self):
        if self.send_thread and self.send_thread.is_alive():
            if not messagebox.askokcancel("退出", "正在发送邮件，确定要退出吗？"):
                return
            self.stop_event.set()
            self.send_thread.join(timeout=2)
        self._save_config()
        self.root.destroy()


def main():
    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    root = tk.Tk()
    try:
        default_font = tkfont.nametofont("TkDefaultFont")
        default_font.configure(size=10, family="Microsoft YaHei UI")
    except Exception:
        pass

    app = MailApp(root)

    def report(exc, val, tb):
        messagebox.showerror("程序错误", "".join(traceback.format_exception(exc, val, tb)))
    root.report_callback_exception = report
    root.mainloop()


if __name__ == "__main__":
    main()
